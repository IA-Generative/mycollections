"""Génère un fichier par format, chacun avec son code témoin FMT-<EXT>-<n>. Contenu fictif.

    uv run --no-project --with python-docx --with python-pptx --with openpyxl --with pillow \\
        python _generer.py <dossier>

Demande LibreOffice (`soffice`) pour les formats convertis. Le .ppt est rangé compressé (.gz) : il
dépasse la taille admise par le dépôt.
"""

import email.message
import gzip
import io
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

OUT = Path(sys.argv[1])
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
FONT = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 34)
TITRE = "Note de contrôle des formats"


def phrase(ext: str, n: int) -> str:
    return f"Le code de contrôle du format {ext.upper()} est FMT-{ext.upper()}-{n}."


def convertir(source: Path, ext: str) -> None:
    """Convertit `source` en `format.<ext>` avec LibreOffice, puis efface la source."""
    sortie = OUT / "_conv"
    subprocess.run([SOFFICE, "--headless", "--convert-to", ext, "--outdir", str(sortie), str(source)], check=True, capture_output=True)
    (sortie / f"{source.stem}.{ext}").rename(OUT / f"format.{ext}")
    source.unlink()


def image(ext: str, n: int) -> Image.Image:
    img = Image.new("RGB", (1400, 260), "white")
    dessin = ImageDraw.Draw(img)
    dessin.text((40, 60), TITRE, fill="black", font=FONT)
    dessin.text((40, 140), phrase(ext, n), fill="black", font=FONT)
    return img


def texte() -> None:
    (OUT / "format.txt").write_text(f"{TITRE}\n\n{phrase('txt', 1101)}\n")
    (OUT / "format.md").write_text(f"# {TITRE}\n\n{phrase('md', 1102)}\n")
    (OUT / "format.csv").write_text(f"rubrique;valeur\ncontrôle;{phrase('csv', 1103)}\n")
    (OUT / "format.html").write_text(f"<html><body><h1>{TITRE}</h1><p>{phrase('html', 1104)}</p></body></html>")
    (OUT / "format.json").write_text(json.dumps({"note": phrase("json", 1105)}, ensure_ascii=False))


def traitement_de_texte() -> None:
    for ext, n in (("docx", 1201), ("doc", 1202), ("odt", 1203), ("rtf", 1204), ("pdf", 1205)):
        doc = Document()
        doc.add_heading(TITRE, 1)
        doc.add_paragraph(phrase(ext, n))
        if ext == "docx":
            doc.save(OUT / "format.docx")
            continue
        source = OUT / f"_src-{ext}.docx"
        doc.save(source)
        convertir(source, ext)


def presentations() -> None:
    for ext, n in (("pptx", 1301), ("ppt", 1302), ("odp", 1303)):
        pres = Presentation()
        diapo = pres.slides.add_slide(pres.slide_layouts[1])
        diapo.shapes.title.text = TITRE
        diapo.placeholders[1].text = phrase(ext, n)
        if ext == "pptx":
            pres.save(OUT / "format.pptx")
            continue
        source = OUT / f"_src-{ext}.pptx"
        pres.save(source)
        convertir(source, ext)


def tableurs() -> None:
    for ext, n in (("xlsx", 1401), ("xls", 1402), ("ods", 1403)):
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.title = "Contrôle"
        feuille.append(["Rubrique", "Valeur"])
        feuille.append(["contrôle", phrase(ext, n)])
        if ext == "xlsx":
            classeur.save(OUT / "format.xlsx")
            continue
        source = OUT / f"_src-{ext}.xlsx"
        classeur.save(source)
        convertir(source, ext)


def autres() -> None:
    message = email.message.EmailMessage()
    message["From"] = "expediteur@fake-domain.name"
    message["To"] = "destinataire@fake-domain.name"
    message["Subject"] = TITRE
    message.set_content(phrase("eml", 1501))
    (OUT / "format.eml").write_bytes(bytes(message))
    image("png", 1601).save(OUT / "format.png")
    image("jpg", 1602).save(OUT / "format.jpg", quality=92)
    image("pdf-scan", 1603).save(OUT / "format-scan.pdf", "PDF", resolution=150)
    image("gif", 1604).save(OUT / "format.gif")
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w") as archive:
        archive.writestr("dedans.txt", phrase("zip", 1701))
    (OUT / "format.zip").write_bytes(tampon.getvalue())


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    texte()
    traitement_de_texte()
    presentations()
    tableurs()
    autres()
    shutil.rmtree(OUT / "_conv", ignore_errors=True)
    (OUT / "format.ppt.gz").write_bytes(gzip.compress((OUT / "format.ppt").read_bytes(), 9))
    (OUT / "format.ppt").unlink()
    for fichier in sorted(OUT.iterdir()):
        print(fichier.name, fichier.stat().st_size)


if __name__ == "__main__":
    main()
